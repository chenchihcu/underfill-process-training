import os
import logging
from typing import Optional
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

logging.basicConfig(level=logging.INFO, format='[%(levelname)s] %(message)s')
log = logging.getLogger(__name__)


def setup_doc(font_name: str = 'Microsoft JhengHei', font_size: int = 10) -> Document:
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = font_name
    style.font.size = Pt(font_size)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    return doc


def add_centered_heading(doc: Document, text: str, level: int = 0) -> None:
    h = doc.add_heading(text, level)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_header_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            t.rows[i].cells[j].text = v


def add_data_table(doc: Document, headers: list[str], data: list[list[str]]) -> None:
    if not headers and not data:
        return
    nrows = len(data) + (1 if headers else 0)
    ncols = len(headers) if headers else (len(data[0]) if data else 0)
    t = doc.add_table(rows=nrows, cols=ncols)
    t.style = 'Table Grid'
    if headers:
        for i, h in enumerate(headers):
            t.rows[0].cells[i].text = h
    start = 1 if headers else 0
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            t.rows[i + start].cells[j].text = val


def add_image(doc: Document, path: str, width: float = 4.5) -> bool:
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(width))
        return True
    log.warning(f'Image not found: {path}')
    return False


def add_image_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style='Normal')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style='List Bullet')
